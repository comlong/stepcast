"""PPT / PDF to video.

Flow:
  upload -> extract each slide's thumbnail, text and speaker notes -> pick slides / how to use the notes / level of detail
  -> one static step per slide, narration from the notes or the AI -> voice-over -> open the editor

Slide image export, in order of preference:
  1. PowerPoint on this computer (COM, highest fidelity)
  2. LibreOffice (convert to PDF, then rasterise)
  3. fallback: draw a simple slide from the text only
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import time
import uuid
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from PIL import Image, ImageDraw

from .. import config, i18n, storage
from ..models import Project, RevealItem, SlideReveal, SlidesCreateReq, Step
from . import clips, slide_reveal

Progress = Optional[Callable[[float, str], None]]

MAX_PAGES = 100
MAX_BYTES = 2 * 1024 ** 3        # largest deck accepted (decks with embedded videos can be big)
RENDER_WIDTH = 1920
THUMB_WIDTH = 360


class SlidesError(RuntimeError):
    pass


def imports_dir() -> Path:
    d = config.DATA_DIR / "_imports"
    d.mkdir(parents=True, exist_ok=True)
    return d


IID_RE = re.compile(r"^imp_[0-9a-f]{10}$")


def import_dir(iid: str) -> Path:
    if not IID_RE.match(iid or ""):
        raise SlidesError(i18n.t("导入记录 id 不合法"))
    return imports_dir() / iid


def load_manifest(iid: str) -> Optional[Dict[str, Any]]:
    if not IID_RE.match(iid or ""):
        return None
    p = import_dir(iid) / "manifest.json"
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding="utf-8"))


def _save_manifest(iid: str, data: Dict[str, Any]) -> None:
    (import_dir(iid) / "manifest.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


# ---- text / notes ----------------------------------------------------------

def _collect_text(shapes, title: str, lines: List[str]) -> None:
    """Collect text recursively: text boxes, tables and shapes inside groups."""
    for shp in shapes:
        try:
            if hasattr(shp, "shapes"):                 # group shape
                _collect_text(shp.shapes, title, lines)
            elif shp.has_text_frame:
                t = shp.text_frame.text.strip()
                if t and t != title:
                    lines.append(t)
            elif getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        except Exception:
            continue


def _notes_text(slide) -> str:
    """Read the speaker notes.

    The standard way is notes_text_frame; but PPTX exported from Google Slides / Keynote sometimes has no body placeholder
    on the notes page and puts the text in an ordinary text box — then fall back to the other text on the notes page (excluding page numbers, headers and footers).
    """
    if not slide.has_notes_slide:
        return ""
    ns = slide.notes_slide
    try:
        tf = ns.notes_text_frame
        if tf is not None and tf.text.strip():
            return tf.text.strip()
    except Exception:
        pass
    skip = {"SLIDE_NUMBER", "HEADER", "FOOTER", "DATE", "SLIDE_IMAGE"}
    parts: List[str] = []
    for shp in ns.shapes:
        try:
            if shp.is_placeholder and shp.placeholder_format.type.name in skip:
                continue
            if shp.has_text_frame and shp.text_frame.text.strip():
                parts.append(shp.text_frame.text.strip())
        except Exception:
            continue
    return "\n".join(parts)


def _pptx_text(pptx_path: Path) -> List[Dict[str, Any]]:
    from pptx import Presentation
    prs = Presentation(str(pptx_path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:
            pass
        lines: List[str] = []
        _collect_text(slide.shapes, title, lines)
        notes = _notes_text(slide)
        hidden = slide._element.get("show") == "0"
        out.append({"i": i, "title": title, "text": "\n".join(lines)[:2000],
                    "notes": notes[:6000], "hidden": hidden})
    return out


# ---- rendering: PowerPoint COM ---------------------------------------------------

def powerpoint_available() -> bool:
    if os.name != "nt" or os.environ.get("VT_DISABLE_POWERPOINT"):   # on servers / in tests never touch the local PowerPoint
        return False
    try:
        import winreg
        winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "PowerPoint.Application").Close()
        import win32com.client  # noqa: F401
        return True
    except Exception:
        return False


def _render_powerpoint(src: Path, out_dir: Path, progress: Progress = None) -> Dict[str, Any]:
    """Export each slide as PNG with the local PowerPoint. Other presentations you have open are never closed."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = None
    pres = None
    was_busy = False
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        try:
            was_busy = bool(app.Visible) or app.Presentations.Count > 0
        except Exception:
            was_busy = True
        pres = app.Presentations.Open(str(src.resolve()), ReadOnly=True, Untitled=False,
                                      WithWindow=False)
        n = int(pres.Slides.Count)
        if n > MAX_PAGES:
            raise SlidesError(i18n.t("最多支持 {max} 页，这个文件有 {n} 页", max=MAX_PAGES, n=n))
        sw, sh = float(pres.PageSetup.SlideWidth), float(pres.PageSetup.SlideHeight)
        w = RENDER_WIDTH
        h = int(round(w * sh / sw))
        hidden: Dict[int, bool] = {}
        for i in range(1, n + 1):
            slide = pres.Slides(i)
            slide.Export(str((out_dir / f"slide_{i:03d}.png").resolve()), "PNG", w, h)
            try:
                hidden[i] = bool(slide.SlideShowTransition.Hidden)
            except Exception:
                hidden[i] = False
            if progress:
                progress(0.1 + 0.8 * i / n, i18n.t("PowerPoint 导出第 {i}/{n} 页", i=i, n=n))
        # old .ppt format: also save a .pptx copy so the notes can be read
        pptx_copy = None
        if src.suffix.lower() == ".ppt":
            pptx_copy = out_dir / "converted.pptx"
            pres.SaveCopyAs(str(pptx_copy.resolve()), 24)   # 24 = ppSaveAsOpenXMLPresentation
        return {"count": n, "hidden": hidden, "engine": "PowerPoint", "pptx": pptx_copy}
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:
            pass
        try:
            if app is not None and not was_busy and app.Presentations.Count == 0:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


# ---- rendering: LibreOffice / PDF ------------------------------------------------

def _soffice() -> Optional[str]:
    for p in (shutil.which("soffice"),
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if p and os.path.exists(p):
            return p
    return None


def _render_pdf(pdf: Path, out_dir: Path, progress: Progress = None,
                p0: float = 0.1, p1: float = 0.9) -> Dict[str, Any]:
    import pymupdf
    doc = pymupdf.open(str(pdf))
    n = doc.page_count
    if n > MAX_PAGES:
        raise SlidesError(i18n.t("最多支持 {max} 页，这个文件有 {n} 页", max=MAX_PAGES, n=n))
    texts: Dict[int, str] = {}
    for i, page in enumerate(doc, 1):
        zoom = RENDER_WIDTH / max(1.0, page.rect.width)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
        pix.save(str(out_dir / f"slide_{i:03d}.png"))
        texts[i] = page.get_text("text").strip()[:2000]
        if progress:
            progress(p0 + (p1 - p0) * i / n, i18n.t("渲染第 {i}/{n} 页", i=i, n=n))
    doc.close()
    return {"count": n, "texts": texts, "engine": "PDF"}


def _render_libreoffice(src: Path, out_dir: Path, progress: Progress = None) -> Dict[str, Any]:
    exe = _soffice()
    if not exe:
        raise SlidesError(i18n.t("没有 LibreOffice"))
    if progress:
        progress(0.1, i18n.t("LibreOffice 转换中…"))
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300,
                   creationflags=0x08000000 if os.name == "nt" else 0)
    pdf = out_dir / (src.stem + ".pdf")
    if not pdf.exists():
        raise SlidesError(i18n.t("LibreOffice 转换失败"))
    r = _render_pdf(pdf, out_dir, progress, 0.4, 0.9)
    r["engine"] = "LibreOffice"
    return r


def _render_text_fallback(slides: List[Dict[str, Any]], out_dir: Path) -> None:
    """Without any renderer, draw a clean simple slide from the title and body text."""
    from .renderer import load_font, wrap_text
    W, H = RENDER_WIDTH, int(RENDER_WIDTH * 9 / 16)
    for s in slides:
        img = Image.new("RGB", (W, H), (255, 255, 255))
        d = ImageDraw.Draw(img)
        d.rectangle((0, 0, W, 14), fill=(255, 92, 57))
        ft, fb = load_font(64, bold=True), load_font(38)
        y = 110
        for ln in wrap_text(d, s["title"] or i18n.t("第 {n} 页", n=s["i"]), ft, W - 240)[:2]:
            d.text((120, y), ln, font=ft, fill=(24, 28, 38))
            y += 86
        y += 30
        for para in (s["text"] or "").splitlines():
            for ln in wrap_text(d, "• " + para, fb, W - 260)[:3]:
                if y > H - 90:
                    break
                d.text((130, y), ln, font=fb, fill=(60, 66, 80))
                y += 56
        img.save(out_dir / f"slide_{s['i']:03d}.png")


# ---- stage 1: upload + parse --------------------------------------------------

def analyze(upload_path: Path, filename: str, progress: Progress = None) -> Dict[str, Any]:
    ext = Path(filename).suffix.lower()
    if ext not in (".pptx", ".ppt", ".pdf"):
        raise SlidesError(i18n.t("只支持 .pptx / .ppt / .pdf"))
    iid = "imp_" + uuid.uuid4().hex[:10]
    d = import_dir(iid)
    d.mkdir(parents=True, exist_ok=True)
    src = d / ("source" + ext)
    shutil.move(str(upload_path), src)

    slides: List[Dict[str, Any]] = []
    engine = ""
    video_source: Optional[Path] = None           # which pptx to take the videos from (for .ppt, the copy converted by PowerPoint)
    if progress:
        progress(0.03, i18n.t("读取文件…"))

    if ext == ".pdf":
        r = _render_pdf(src, d, progress)
        engine = r["engine"]
        slides = [{"i": i, "title": (r["texts"][i].splitlines() or [""])[0][:80],
                   "text": r["texts"][i], "notes": "", "hidden": False}
                  for i in range(1, r["count"] + 1)]
    else:
        pptx_for_text = src if ext == ".pptx" else None
        rendered = None
        if powerpoint_available():
            try:
                rendered = _render_powerpoint(src, d, progress)
                engine = "PowerPoint"
                if rendered.get("pptx"):
                    pptx_for_text = rendered["pptx"]
            except SlidesError:
                raise
            except Exception as e:
                if progress:
                    progress(0.1, i18n.t("PowerPoint 导出失败，改用其他方式：{error}", error=str(e)[:80]))
                rendered = None
        if pptx_for_text is None:
            raise SlidesError(i18n.t(".ppt 老格式需要本机安装 PowerPoint，或先另存为 .pptx"))
        slides = _pptx_text(pptx_for_text)
        video_source = pptx_for_text
        if len(slides) > MAX_PAGES:
            raise SlidesError(i18n.t("最多支持 {max} 页，这个文件有 {n} 页", max=MAX_PAGES, n=len(slides)))
        if rendered:
            for s in slides:
                s["hidden"] = s["hidden"] or rendered["hidden"].get(s["i"], False)
        else:
            try:
                _render_libreoffice(src, d, progress)
                engine = "LibreOffice"
            except Exception:
                _render_text_fallback(slides, d)
                engine = i18n.t("简版（未检测到 PowerPoint / LibreOffice）")

    for s in slides:
        png = d / f"slide_{s['i']:03d}.png"
        if not png.exists():
            _render_text_fallback([s], d)
        with Image.open(png) as im:
            s["w"], s["h"] = im.size
            th = im.convert("RGB")
            th.thumbnail((THUMB_WIDTH, THUMB_WIDTH))
            th.save(d / f"thumb_{s['i']:03d}.jpg", "JPEG", quality=82)

    # videos on the slides: an exported image keeps only the poster, so extract the video files themselves here
    videos: Dict[int, List[Dict[str, Any]]] = {}
    if video_source is not None:
        if progress:
            progress(0.93, i18n.t("查找幻灯片里的视频…"))
        try:
            videos = clips.pptx_videos(video_source, d)
        except Exception:
            videos = {}
    for s in slides:
        s["videos"] = videos.get(s["i"], [])

    manifest = {
        "id": iid, "filename": filename, "ext": ext, "engine": engine,
        "created_at": time.time(), "count": len(slides),
        "with_notes": sum(1 for s in slides if s["notes"]),
        "with_videos": sum(1 for s in slides if s["videos"]),
        "slides": slides,
    }
    _save_manifest(iid, manifest)
    if progress:
        progress(1.0, i18n.t("解析完成：{n} 页，{notes} 页有备注（{engine}）", n=len(slides),
                             notes=manifest["with_notes"], engine=engine))
    return manifest


# ---- stage 2: create the project -----------------------------------------------------

def create_project(iid: str, req: SlidesCreateReq, progress: Progress = None) -> Dict[str, Any]:
    from . import script_gen, tts

    man = load_manifest(iid)
    if not man:
        raise SlidesError(i18n.t("导入记录不存在或已过期，请重新上传"))
    chosen = set(req.selected) if req.selected else {s["i"] for s in man["slides"] if not s["hidden"]}
    slides = [s for s in man["slides"] if s["i"] in chosen]
    if not slides:
        raise SlidesError(i18n.t("没有选中任何页"))

    name = req.name or Path(man["filename"]).stem
    proj = storage.create(name, req.language)
    proj.source = "slides"
    proj.voice = req.voice or proj.voice      # none chosen: use the voice from the settings (storage.create already picked one for the narration language)
    proj.settings = {
        "slides_notes_mode": req.notes_mode,
        "slides_missing": req.missing,
        "slides_detail": req.detail,
        "zoom_enabled": False,
        "dim_background": False,
        "show_cursor": False,
        "show_step_badge": False,
        "browser_frame": False,
    }
    if req.dialogue:
        # two-person Q&A: the host asks, the expert explains. Intro, outro and narration without lines are read by the host
        from . import dialogue
        proj.settings["dialogue"] = True
        proj.speakers = dialogue.default_speakers(req.language or proj.language, req.host_voice, req.expert_voice)
        proj.voice = proj.speakers[0].voice
    src_dir = import_dir(iid)
    warnings: List[str] = []
    reveal = _export_reveal(man, slides, src_dir, req, warnings, progress)
    proj.settings["slides_reveal"] = bool(reveal)
    for s in slides:
        step = Step(kind="slide", page_title=s["title"][:120], img_w=s.get("w", 0),
                    img_h=s.get("h", 0), title="", zoom=False, highlight=False)
        step.slide_text = s["text"]
        step.slide_notes = s["notes"]
        fname = f"{step.id}.png"
        shutil.copyfile(src_dir / f"slide_{s['i']:03d}.png", storage.screenshots_dir(proj.id) / fname)
        step.screenshot = fname
        step.viewport_w, step.viewport_h = step.img_w, step.img_h
        if s["i"] in reveal:
            step.reveal = _attach_reveal(proj, step, reveal[s["i"]], src_dir / "reveal")
            # the cover appears as a whole by default: its big title is often in the middle of the page, not in the title area, and would be split into several items. The data is still generated;
            # to reveal it one by one, switch it on for this slide in the editor
            # (slides with the author's own animations are revealed as animated, cover or not)
            step.reveal.enabled = step.reveal.mode == "timeline" or not _is_cover(s, man)
            if req.videos and s.get("videos") and any(it.after_media for it in step.reveal.items):
                step.sequence = "video_first"        # the author made the text come in after the video has played
        step.index = len(proj.steps)
        proj.steps.append(step)
        if req.videos:
            for v in s.get("videos") or []:
                proj.steps.append(_video_step(proj, s, v, src_dir))
    storage.save(proj)
    if progress:
        progress(0.1, i18n.t("已创建项目，共 {n} 页", n=len(slides)))

    from .llm import LLMError
    warning = ""
    try:
        res = script_gen.generate_slides_script(
            proj, notes_mode=req.notes_mode, detail=req.detail, missing=req.missing,
            progress=lambda f, m: progress(0.1 + f * 0.5, m) if progress else None)
    except LLMError as e:
        # no AI key or the AI service failed: the project already exists, so don't fail the whole import (this used to leave a half-made project and not open the editor).
        # narration already taken from the notes is kept; the rest stays empty for now
        res = None
        warning = i18n.t("解说没能自动生成：{error}\n项目已经建好，可以在编辑器里自己写解说，或者配好 AI 后点「① 生成解说」。",
                         error=e)
    storage.save(proj)

    voiced = None
    if req.auto_voice:
        voiced = tts.synth_project(
            proj, voice=proj.voice,
            progress=lambda f, m: progress(0.6 + f * 0.38, m) if progress else None)
        storage.save(proj)

    if progress:
        progress(1.0, i18n.t("完成，进入编辑器"))
    warning = "\n\n".join(w for w in [warning, *warnings] if w)
    return {"project_id": proj.id, "slides": len(slides), "script": res, "tts": voiced, "warning": warning}


def _is_cover(s: Dict[str, Any], man: Dict[str, Any]) -> bool:
    """Cover = slide 1 of a multi-slide deck with little text (title, subtitle, date). A first slide that is already content doesn't count."""
    return s["i"] == 1 and man.get("count", 0) > 1 and len(re.sub(r"\s+", "", s.get("text") or "")) < 120


def _export_reveal(man: Dict[str, Any], slides: List[Dict[str, Any]], src_dir: Path, req: SlidesCreateReq,
                   warnings: List[str], progress: Progress) -> Dict[int, Dict[str, Any]]:
    """Reveal one by one: export base images and items for the selected slides with PowerPoint. Only for slides exported by PowerPoint
    (base image, items and full slide must come from the same renderer for fonts and positions to match)."""
    if not req.reveal or man.get("ext") not in (".pptx", ".ppt"):
        return {}
    if man.get("engine") != "PowerPoint" or not powerpoint_available():
        warnings.append(i18n.t("逐条出现需要本机的 PowerPoint，这次没有生成，所有页整页一起出现。"))
        return {}
    pages = [s["i"] for s in slides]
    try:
        return slide_reveal.export(
            src_dir / ("source" + man["ext"]), pages, src_dir / "reveal", RENDER_WIDTH,
            progress=lambda f, m: progress(0.02 + f * 0.08, m) if progress else None,
            label=lambda i, n: i18n.t("分析页面内容，准备逐条出现 {i}/{n}", i=i, n=n))
    except Exception as e:  # noqa: BLE001 — revealing is a nice extra; on errors import as usual
        warnings.append(i18n.t("逐条出现没能生成（{error}），所有页整页一起出现。", error=str(e)[:160]))
        return {}


def _attach_reveal(proj: Project, step: Step, meta: Dict[str, Any], src: Path) -> SlideReveal:
    """Copy this slide's base image and item images into the project; file names follow the step."""
    shots = storage.screenshots_dir(proj.id)
    clean = f"{step.id}_clean.png"
    shutil.copyfile(src / meta["clean"], shots / clean)
    items = []
    for k, it in enumerate(meta["items"], 1):
        name = f"{step.id}_r{k:02d}.png"
        shutil.copyfile(src / it["file"], shots / name)
        timing = {f: it[f] for f in ("beat", "offset", "anim", "exit_beat", "exit_offset", "after_media", "media", "dur", "side",
                                      "motion_beat", "motion_offset", "motion_dur", "motion_path", "motion_accel", "motion_decel") if f in it}
        items.append(RevealItem(file=name, x=int(it["x"]), y=int(it["y"]), text=it.get("text", ""), **timing))
    return SlideReveal(clean=clean, items=items, mode=meta.get("mode", ""))


def _video_step(proj: Project, s: Dict[str, Any], v: Dict[str, Any], src_dir: Path) -> Step:
    """A video on a slide -> a "video" step right after that slide. Its base image is the slide (the video plays in place)."""
    step = Step(kind="video", page_title=s["title"][:120], img_w=s.get("w", 0), img_h=s.get("h", 0),
                title="", zoom=False, highlight=False)
    fname = f"{step.id}.png"
    shutil.copyfile(src_dir / f"slide_{s['i']:03d}.png", storage.screenshots_dir(proj.id) / fname)
    step.screenshot = fname
    step.viewport_w, step.viewport_h = step.img_w, step.img_h
    step.clip = clips.clip_from_import(v)
    if v.get("file") and not v.get("missing"):
        name = f"{step.id}{Path(v['file']).suffix}"
        clips.media_dir(proj.id).mkdir(parents=True, exist_ok=True)
        # videos can be hundreds of MB: the import folder and the project are on the same drive, so hard-link instead of copying
        clips.link_or_copy(src_dir / v["file"], clips.media_dir(proj.id) / name)
        step.clip.file = name
    else:
        # linked / online / broken videos have no file: leave them out of the video for now (otherwise the slide would appear twice); added automatically after upload
        step.include = False
    step.index = len(proj.steps)
    return step


def cleanup_old_imports(max_age_hours: float = 48) -> int:
    n = 0
    now = time.time()
    for d in imports_dir().iterdir():
        try:
            if d.is_dir() and now - d.stat().st_mtime > max_age_hours * 3600:
                shutil.rmtree(d, ignore_errors=True)
                n += 1
        except Exception:
            pass
    return n
