"""Test decks: an embedded video (with PowerPoint trim settings), a video linked to a local file, an online video, a video inside a group;
plus a video with speech (for testing "speech to subtitles")."""
import copy, subprocess, sys
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

VID = Path(sys.argv[1]) / "vid"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def slide(title):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    return s


s1 = slide("第 1 页")
s2 = slide("第 2 页：嵌入视频（剪辑 2~7 秒）")
mv = s2.shapes.add_movie(str(VID / "clip10.mp4"), Inches(1), Inches(2), Inches(6.4), Inches(3.6), mime_type="video/mp4")
media = mv._element.find(f".//{{{P14}}}media")
etree.SubElement(media, f"{{{P14}}}trim", st="2000", end="3000")      # trim 2 seconds from the start and 3 seconds from the end

# linked to a file on the author's computer: turn the embedded video into an external link
s3 = slide("第 3 页：链接的视频")
mv3 = s3.shapes.add_movie(str(VID / "clip10.mp4"), Inches(2), Inches(2), Inches(6.4), Inches(3.6), mime_type="video/mp4")
rid = s3.part.relate_to("file:///C:/Users/author/Videos/product-demo.mp4", RT.VIDEO, is_external=True)
mv3._element.find(".//" + qn("a:videoFile")).set(qn("r:link"), rid)
ext = mv3._element.find(".//" + qn("p:extLst"))
ext.getparent().remove(ext)

# online video
s4 = slide("第 4 页：在线视频")
mv4 = s4.shapes.add_movie(str(VID / "clip10.mp4"), Inches(2), Inches(2), Inches(6.4), Inches(3.6), mime_type="video/mp4")
rid = s4.part.relate_to("https://www.youtube.com/embed/dQw4w9WgXcQ", RT.VIDEO, is_external=True)
mv4._element.find(".//" + qn("a:videoFile")).set(qn("r:link"), rid)
ext = mv4._element.find(".//" + qn("p:extLst"))
ext.getparent().remove(ext)

# video inside a group: the group is moved to the bottom right as a whole, and the video's coordinates must be converted
s5 = slide("第 5 页：组合里的视频")
grp = s5.shapes.add_group_shape()
mv5 = s5.shapes.add_movie(str(VID / "clip10.mp4"), 0, 0, Inches(8), Inches(4.5), mime_type="video/mp4")
grp._element.append(mv5._element)                 # move into the group (group coordinates 0,0 / 8x4.5 inches)
x = grp._element.grpSpPr.get_or_add_xfrm()
x.set  # noqa
from lxml import etree as ET
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
for tag in ("off", "ext", "chOff", "chExt"):
    for old_el in x.findall(f"{{{A}}}{tag}"):
        x.remove(old_el)
# the group on the page: top left (8, 4.5) inches, size 4x2.25 inches — its content is scaled to half
ET.SubElement(x, f"{{{A}}}off", x=str(Inches(8)), y=str(Inches(4.5)))
ET.SubElement(x, f"{{{A}}}ext", cx=str(Inches(4)), cy=str(Inches(2.25)))
ET.SubElement(x, f"{{{A}}}chOff", x="0", y="0")
ET.SubElement(x, f"{{{A}}}chExt", cx=str(Inches(8)), cy=str(Inches(4.5)))
prs.save(VID / "deck_cases.pptx")

# video with speech: solid color frames + line0.mp3 (one Chinese sentence)
subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                "color=c=navy:size=640x360:rate=25", "-i", str(Path(sys.argv[1]) / "line0.mp3"),
                "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", str(VID / "speech.mp4")],
               check=True)
print("ok")
