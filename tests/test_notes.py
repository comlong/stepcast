"""PPT speaker notes as narration."""
import copy
import json
import os
import shutil
import sys
import time
from pathlib import Path

SP = Path(sys.argv[1])
DATA = SP / "notes_data"
shutil.rmtree(DATA, ignore_errors=True)
DATA.mkdir(parents=True)
os.environ["VT_DATA_DIR"] = str(DATA)
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from backend import main, storage  # noqa: E402
from backend.services import script_gen, slides, tts  # noqa: E402

slides.powerpoint_available = lambda: False      # never touch the PowerPoint you have open
tts.synth = lambda text, voice, out, *a, **k: (out.write_bytes(b"\0" * 2048), (1.0, []))[1]

c = TestClient(main.app, base_url="http://127.0.0.1:8756", headers={"Origin": "http://127.0.0.1:8756"})
fails = []


def check(name, cond, detail=""):
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  —— {detail}" if detail else ""), flush=True)
    if not cond:
        fails.append(name)


class RecordingLLM:
    calls = []

    def __init__(self, *a, **k):
        pass

    def chat_json(self, messages, **k):
        prompt = messages[-1]["content"]
        RecordingLLM.calls.append(prompt)
        import re
        idx = [int(x) for x in re.findall(r'"i": (\d+)', prompt)]
        return {"title": "AI 标题", "intro": "AI 片头。", "outro": "AI 片尾。",
                "steps": [{"i": i, "title": f"页{i}", "narration": f"AI 写的第 {i} 页。"} for i in idx]}


script_gen.get_client = RecordingLLM


def wait(j):
    while True:
        r = c.get(f"/api/jobs/{j['id']}").json()
        if r["status"] not in ("pending", "running"):
            return r
        time.sleep(0.2)


# ---------------------------------------------------------------- build a deck
NOTES = {1: "欢迎来到销售对话培训，这是我们认为标准的开场讲解。",
         2: "第二页，先确认客户需求，再介绍产品。\n注意语气要自然。",
         4: "第四页的备注写在普通文本框里，不在标准占位符中。"}
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
for i in range(1, 5):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = f"第 {i} 页标题"
    s.placeholders[1].text_frame.text = f"第 {i} 页的正文要点"
    if i in (1, 2):
        s.notes_slide.notes_text_frame.text = NOTES[i]
    if i == 4:
        ns = s.notes_slide
        body = ns.notes_placeholder
        body._element.getparent().remove(body._element)          # remove the standard notes placeholder
        # notes pages can't take a text box directly: create one on the slide, then move its XML to the notes page
        tb = s.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(2))
        tb.text_frame.text = NOTES[4]
        el = tb._element
        el.getparent().remove(el)
        ns.shapes._spTree.append(el)
deck = SP / "notes_deck.pptx"
prs.save(deck)


def upload(path, name):
    with open(path, "rb") as f:
        r = wait(c.post("/api/import/slides", files={"file": (name, f)}).json())
    assert r["status"] == "done", r
    return r["result"]


print("\n== 解析备注 ==")
man = upload(deck, "销售对话.pptx")
got = {s["i"]: s["notes"] for s in man["slides"]}
check("标准占位符里的备注读到了", got[1] == NOTES[1] and got[2] == NOTES[2])
check("第 3 页没有备注", got[3] == "")
check("写在普通文本框里的备注也读到了（兜底）", got[4] == NOTES[4], repr(got[4]))
check("统计出 3 页有备注", man["with_notes"] == 3, man["with_notes"])

print("\n== 备注原文 + 没备注的页留空 ==")
RecordingLLM.calls.clear()
r = wait(c.post(f"/api/import/slides/{man['id']}/create",
                json={"notes_mode": "verbatim", "missing": "empty", "language": "zh-CN"}).json())
p = storage.load(r["result"]["project_id"])
narr = {s.index + 1: s.narration for s in p.steps}
check("任务成功", r["status"] == "done", r.get("error", ""))
check("完全没有调用 DeepSeek", not RecordingLLM.calls, len(RecordingLLM.calls))
check("有备注的页 = 备注原文（换行已并成一段）",
      narr[1] == NOTES[1] and narr[2] == NOTES[2].replace("\n", "") and narr[4] == NOTES[4], narr)
check("没备注的第 3 页留空", narr[3] == "")
check("选择被记在项目里", p.settings.get("slides_notes_mode") == "verbatim"
      and p.settings.get("slides_missing") == "empty")
check("有备注的页已配音、空页没有", [bool(s.audio) for s in p.steps] == [True, True, False, True],
      [bool(s.audio) for s in p.steps])
PID_EMPTY = p.id

print("\n== 备注原文 + 没备注的页 AI 补写 ==")
RecordingLLM.calls.clear()
r = wait(c.post(f"/api/import/slides/{man['id']}/create",
                json={"notes_mode": "verbatim", "missing": "ai", "language": "zh-CN"}).json())
p = storage.load(r["result"]["project_id"])
narr = {s.index + 1: s.narration for s in p.steps}
sent = "\n".join(RecordingLLM.calls)
check("只调用一次 DeepSeek", len(RecordingLLM.calls) == 1, len(RecordingLLM.calls))
check("请求里只有第 3 页", "第 3 页标题" in sent and "第 1 页标题" not in sent and "第 4 页标题" not in sent)
check("备注原文没有被发出去", not any(n.split("，")[0] in sent for n in NOTES.values()))
check("第 3 页由 AI 补写，其余是备注原文",
      narr[3].startswith("AI 写的") and narr[1] == NOTES[1] and narr[4] == NOTES[4], narr)

print("\n== AI 参考备注改写（对照）==")
RecordingLLM.calls.clear()
r = wait(c.post(f"/api/import/slides/{man['id']}/create",
                json={"notes_mode": "reference", "language": "zh-CN"}).json())
p_ref = storage.load(r["result"]["project_id"])
sent = "\n".join(RecordingLLM.calls)
check("参考模式会把备注发给 AI", "标准的开场讲解" in sent)
check("参考模式解说是 AI 写的", all(s.narration.startswith("AI 写的") for s in p_ref.steps))

print("\n== 已导入的项目，在「生成解说」里改用备注原文 ==")
RecordingLLM.calls.clear()
r = wait(c.post(f"/api/projects/{p_ref.id}/script",
                json={"notes_mode": "verbatim", "missing": "empty", "overwrite": True}).json())
p2 = storage.load(p_ref.id)
narr = {s.index + 1: s.narration for s in p2.steps}
check("任务成功", r["status"] == "done", r.get("error", ""))
check("没有调用 DeepSeek", not RecordingLLM.calls)
check("有备注的页换成了备注原文", narr[1] == NOTES[1] and narr[4] == NOTES[4], narr)
check("没备注的页保留原来 AI 写的内容（留空模式不删已有文字）", narr[3].startswith("AI 写的"))
check("项目记住了新的选择", p2.settings.get("slides_notes_mode") == "verbatim")

print("\n== 之后点「一键生成」沿用这个选择 ==")
RecordingLLM.calls.clear()
c.patch(f"/api/projects/{p_ref.id}/steps/{p2.steps[0].id}", json={"narration": "", "caption": ""})
c.patch(f"/api/projects/{p_ref.id}/steps/{p2.steps[0].id}", json={"slide_notes": "我在编辑器里改过的备注。"})
import backend.services.video as video
video.render_project = lambda proj, **k: {"file": ""}
r = wait(c.post(f"/api/projects/{p_ref.id}/auto", json={}).json())
p3 = storage.load(p_ref.id)
check("一键生成成功", r["status"] == "done", r.get("error", ""))
check("第 1 页用了编辑器里改过的备注", p3.steps[0].narration == "我在编辑器里改过的备注。", p3.steps[0].narration)
check("没有调用 DeepSeek", not RecordingLLM.calls)

print("\n== 备注变成解说：折行接上，段落和分点各自成句 ==")
N = script_gen.notes_as_narration
check("小标题 + 空行 + 正文：小标题自己成一句（配音会停顿、字幕在这里断开），不和下一句粘在一起",
      N("Opening Remarks\n\nA good way to start is a question.") == "Opening Remarks. A good way to start is a question.",
      N("Opening Remarks\n\nA good way to start is a question."))
check("一句话中间的折行：直接接上，不加句号", N("One sentence that wraps\nin the middle and ends.") == "One sentence that wraps in the middle and ends.")
check("分点：每一项成一句，项目符号不念出来，冒号后面不再加句号",
      N("Key points:\n- Fast charging\n- Long range\n• Quiet cabin") == "Key points: Fast charging. Long range. Quiet cabin",
      N("Key points:\n- Fast charging\n- Long range\n• Quiet cabin"))
check("编号保留", N("1. First step\n2. Second step") == "1. First step. 2. Second step", N("1. First step\n2. Second step"))
check("中文：直接拼接，补的是「。」，已有句号的不重复加",
      N("开场白\n\n先确认客户需求，\n再介绍产品。\n注意语气要自然。") == "开场白。先确认客户需求，再介绍产品。注意语气要自然。",
      N("开场白\n\n先确认客户需求，\n再介绍产品。\n注意语气要自然。"))
check("段落以逗号结尾：句子还没完，不加句号", N("Heading,\n\ncontinues after a comma") == "Heading, continues after a comma")
check("编辑器「用备注替换解说词」用的是同一套规则", c.post("/api/notes/narration", json={"text": "Opening Remarks\n\n- One\n- Two"}).json()
      == {"text": "Opening Remarks. One. Two"})

print("\n== 导入后自动一键生成（开关默认关闭） ==")
rendered = []
video.render_project = lambda proj, **k: (rendered.append(proj.id), {"file": "x.mp4"})[1]
RecordingLLM.calls.clear()
r = wait(c.post(f"/api/import/slides/{man['id']}/create",
                json={"notes_mode": "verbatim", "missing": "empty", "language": "zh-CN", "auto_generate": True}).json())
res = r["result"]
check("打开开关：导入完成后接着启动了一键生成", r["status"] == "done" and res.get("auto_job"), res)
a = wait({"id": res["auto_job"]}) if res.get("auto_job") else {}
pa = storage.load(res["project_id"])
check("一键生成做完：配音、渲染都做了", a.get("status") == "done" and rendered == [pa.id], (a.get("status"), a.get("error"), rendered))
check("解说按导入时的选择来：没备注的页仍然留空，不让 AI 补写", pa.steps[2].narration == "" and not RecordingLLM.calls
      and "script" not in (a.get("result") or {}), (pa.steps[2].narration, len(RecordingLLM.calls)))
check("有解说的页都配了音", [bool(s.audio) for s in pa.steps] == [True, True, False, True], [bool(s.audio) for s in pa.steps])
check("任务进度从配音开始，不显示「第 1/3 步：生成解说脚本」", not any("生成解说脚本" in str(x) for x in a.get("log", [])), a.get("log", [])[:3])

rendered.clear()
r = wait(c.post(f"/api/import/slides/{man['id']}/create",
                json={"notes_mode": "verbatim", "missing": "empty", "language": "zh-CN"}).json())
check("开关关着（默认）：只导入，不自动渲染", r["status"] == "done" and not r["result"].get("auto_job") and not rendered, r["result"].get("auto_job"))

rendered.clear()
r = wait(c.post(f"/api/import/slides/{man['id']}/create",
                json={"notes_mode": "verbatim", "missing": "empty", "language": "zh-CN", "auto_voice": False, "auto_generate": True}).json())
a = wait({"id": r["result"]["auto_job"]})
pa = storage.load(r["result"]["project_id"])
check("没勾「写完解说自动配音」也一样：一键生成里会配音", a["status"] == "done" and [bool(s.audio) for s in pa.steps] == [True, True, False, True]
      and rendered == [pa.id], [bool(s.audio) for s in pa.steps])


class NoAI:
    def __init__(self, *a, **k):
        from backend.services.llm import LLMError
        raise LLMError("test: no AI configured")


script_gen.get_client = NoAI
rendered.clear()
r = wait(c.post(f"/api/import/slides/{man['id']}/create",
                json={"notes_mode": "ignore", "language": "zh-CN", "auto_generate": True}).json())
script_gen.get_client = RecordingLLM
check("要 AI 写解说但 AI 用不了：项目照样建好，但不自动配音渲染，并说明原因",
      r["status"] == "done" and not r["result"].get("auto_job") and not rendered and "没有自动接着配音和渲染" in r["result"].get("warning", ""),
      (r["result"].get("auto_job"), r["result"].get("warning", "")[-60:]))

print("\n== PDF ==")
pdf = SP / "pptspike" / "export_slides.pdf"
man_pdf = upload(pdf, "导出.pdf")
check("PDF 解析成功", man_pdf["count"] == 3)
check("PDF 没有备注", man_pdf["with_notes"] == 0 and man_pdf["ext"] == ".pdf")

print("\n" + ("全部通过" if not fails else f"失败 {len(fails)} 项：{fails}"))
print("IIDS", json.dumps({"pptx": man["id"], "pdf": man_pdf["id"]}))
