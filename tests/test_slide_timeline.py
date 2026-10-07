"""Slide animations: how PowerPoint's effect list becomes beats and layers (pure logic, no PowerPoint needed).

The effect lists below are what PowerPoint itself wrote for decks built in tests/test_anim_ppt.py (read back through its object model):
"fade on click, text with it, a badge right after", "bullets built by paragraph", "an element that exits again", "a video, then callouts"."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from backend.models import RevealItem  # noqa: E402
from backend.services import slide_timeline as T  # noqa: E402
from backend.services.slide_timeline import AFTER, CLICK, WITH, Effect, ShapeInfo  # noqa: E402

fails = []


def check(name, cond, detail=""):
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  —— {detail}" if detail != "" else ""), flush=True)
    cond or fails.append(name)


FADE, APPEAR, FLY, MEDIA = 10, 1, 2, 83

print("\n== 1. 分类（PowerPoint 实测的编号） ==")
check("入场 1-51、53；退出同号加 Exit", T.classify(10, False) == "in" and T.classify(53, False) == "in"
      and T.classify(10, True) == "out" and T.classify(2, True) == "out")
check("强调 52、54-82 不算", all(T.classify(i, False) == "other" for i in (52, 54, 60, 82)))
check("路径 86 起是动作路径（光点沿线移动）", T.classify(86, False) == "path" and T.classify(149, False) == "path")
check("播放视频 83；暂停、停止不是", T.classify(83, False) == "media" and T.classify(84, False) == "other"
      and T.classify(85, False) == "other")

print("\n== 2. 时间：点击、与上一项同时、在上一项之后 ==")
eff = [Effect(2, FADE, trigger=CLICK, dur=0.5), Effect(3, FADE, trigger=WITH, dur=0.5),
       Effect(4, APPEAR, trigger=AFTER, dur=0.001)]
tm = T.schedule(eff)
check("同时开始的从 0 起", tm[0].start == 0 and tm[1].start == 0)
check("「在上一项之后」排在前面最长的那项结束后（PowerPoint 写的是 delay=500）", abs(tm[2].start - 0.5) < 1e-9 and tm[2].group == 0, tm[2].start)
tm = T.schedule([Effect(1, FADE, trigger=CLICK), Effect(2, FADE, trigger=CLICK), Effect(3, FADE, trigger=WITH)])
check("每次点击是一个新的组，同组的「同时」跟着它", [t.group for t in tm] == [0, 1, 1])
tm = T.schedule([Effect(1, MEDIA, trigger=AFTER, dur=10.0), Effect(2, FADE, trigger=AFTER)])
check("视频的时长就是视频的长度：后面「在上一项之后」的从第 10 秒起（PowerPoint 写的是 delay=10000）",
      abs(tm[1].start - 10.0) < 1e-9, tm[1].start)
tm = T.schedule([Effect(1, FADE, trigger=CLICK, dur=0.5), Effect(2, FADE, trigger=AFTER, delay=1.5, dur=0.5)])
check("延迟从那一步的起点算", abs(tm[1].start - 2.0) < 1e-9, tm[1].start)

print("\n== 3. 一页：卡片 + 文字 + 角标 + 按段落出现的要点 + 会消失的提示 ==")
# shapes in z-order: 0 panel, 1 title (static) | 2 card, 3 its text, 4 badge | 5 bullets (4 paragraphs) | 6 note
shapes = [ShapeInfo("panel"), ShapeInfo("Title"), ShapeInfo(""), ShapeInfo("Card one text"), ShapeInfo("1"),
          ShapeInfo("First point\nSecond point\nsub of second\nThird point",
                    ["First point", "Second point", "sub of second", "Third point"]), ShapeInfo("Gone later")]
eff = [Effect(2, FADE, trigger=CLICK, dur=0.5), Effect(3, FADE, trigger=WITH, dur=0.5, background=False),
       Effect(4, APPEAR, trigger=AFTER, dur=0.001),
       Effect(5, FLY, trigger=CLICK, para=1, background=False), Effect(5, FLY, trigger=CLICK, para=2, background=False),
       Effect(5, FLY, trigger=WITH, para=3, background=False), Effect(5, FLY, trigger=CLICK, para=4, background=False),
       Effect(6, FADE, trigger=CLICK, dur=0.5), Effect(6, FADE, exit=True, trigger=AFTER, delay=1.5, dur=0.5)]
a = T.analyze(eff, shapes)
check("有动画的页能分析", a is not None and a.clicks and not a.media)
check("5 个节拍：卡片组、要点 1、要点 2+3、要点 4、提示", a.beats == 5, a.beats)
L = a.layers
check("层按 PPT 的上下顺序：第一层是底板和标题（没有动画，合在一起）", L[0].shapes == [0, 1] and not L[0].animated, L[0])
check("后面是卡片、文字、角标", [ly.shapes for ly in L[1:4]] == [[2], [3], [4]])
check("卡片、文字、角标在第 1 步；角标晚 0.5 秒（接在淡入之后）",
      [ly.enter[0] for ly in L[1:4]] == [0, 0, 0] and abs(L[3].enter[1] - 0.5) < 1e-9 and L[1].enter[1] == 0, [ly.enter for ly in L[1:4]])
check("「出现」的角标是瞬间出现，其他淡入", L[3].enter[2] == "appear" and L[1].enter[2] == "fade")
bul = [ly for ly in L if ly.shapes == [5]]
check("按段落的文本框：先是文本框本身（没有文字），再每段一层", bul[0].paras == [] and [ly.paras for ly in bul[1:]] == [[1], [2], [3], [4]], [ly.paras for ly in bul])
check("文本框本身不跟着第一段出现（第一项没勾「动画附带的形状」）", not bul[0].animated)
check("要点各自的步骤；子要点（与上一项同时）和上一个要点同一步",
      [ly.enter[0] for ly in bul[1:]] == [1, 2, 2, 3] and bul[3].enter[1] == 0, [ly.enter for ly in bul[1:]])
check("每个要点的文字", [ly.text for ly in bul[1:]] == ["First point", "Second point", "sub of second", "Third point"])
note = L[-1]
check("会消失的提示：第 5 步出现，同一步里 2.0 秒后消失（0.5 的淡入 + 1.5 的延迟）",
      note.enter[0] == 4 and note.exit == (4, 2.0), (note.enter, note.exit))
check("每步说的内容（给解说对齐用）", a.beat_texts() == ["Card one text\n1", "First point", "Second point\nsub of second", "Third point", "Gone later"],
      a.beat_texts())

print("\n== 4. 后面还有文字的情况 ==")
a2 = T.analyze([Effect(0, FADE, trigger=CLICK)], [ShapeInfo("A"), ShapeInfo("B")])
check("只有一个有动画的形状：另一个是静态层", [ly.shapes for ly in a2.layers] == [[0], [1]]
      and [ly.animated for ly in a2.layers] == [True, False])
shapes_b = [ShapeInfo("under"), ShapeInfo("card"), ShapeInfo("over")]
a3 = T.analyze([Effect(1, FADE, trigger=CLICK)], shapes_b)
check("动画形状上面的静态形状是单独一层，仍然盖在它上面（不会被后出现的卡片盖住）",
      [ly.shapes for ly in a3.layers] == [[0], [1], [2]] and [ly.animated for ly in a3.layers] == [False, True, False], [ly.shapes for ly in a3.layers])
a4 = T.analyze([Effect(0, FADE, exit=True, trigger=CLICK)], [ShapeInfo("x")])
check("只有退出动画：不拆开", a4 is None)
a5 = T.analyze([Effect(0, 61, trigger=CLICK)], [ShapeInfo("x")])
check("只有强调动画：不拆开", a5 is None)
a6 = T.analyze([], [ShapeInfo("x")])
check("没有动画：不拆开", a6 is None)
a7 = T.analyze([Effect(1, FADE, trigger=CLICK)], [ShapeInfo("hidden", visible=False), ShapeInfo("card")])
check("作者隐藏的形状不进来", [ly.shapes for ly in a7.layers] == [[1]])

print("\n== 5. 没有点击：自己跑的动画，每个「在上一项之后」是一步 ==")
auto = [Effect(0, FADE, trigger=AFTER, dur=0.5), Effect(1, FADE, trigger=WITH, dur=0.5),
        Effect(2, FADE, trigger=AFTER, dur=0.5), Effect(3, FADE, trigger=AFTER, delay=0.2, dur=0.5)]
a8 = T.analyze(auto, [ShapeInfo("t0"), ShapeInfo("t1"), ShapeInfo("t2"), ShapeInfo("t3")])
check("没有点击", not a8.clicks)
check("三步：第一、二个一起，然后各一步", a8.beats == 3 and [ly.enter[0] for ly in a8.layers] == [0, 0, 1, 2], [ly.enter for ly in a8.layers])
check("延迟留在步骤里", abs(a8.layers[3].enter[1]) < 1e-9, a8.layers[3].enter)

print("\n== 6. 只让东西消失的步骤并到相邻的一步 ==")
ex = [Effect(0, FADE, trigger=CLICK), Effect(0, FADE, exit=True, trigger=CLICK), Effect(1, FADE, trigger=CLICK)]
a9 = T.analyze(ex, [ShapeInfo("one"), ShapeInfo("two")])
check("第二次点击只是让第一个消失：并到下一步（第三次点击）", a9.beats == 2 and a9.layers[0].exit[0] == 1 and a9.layers[1].enter[0] == 1,
      (a9.beats, a9.layers[0].exit, a9.layers[1].enter))

print("\n== 7. 视频，然后提示 ==")
vshapes = [ShapeInfo("", media=True), ShapeInfo("Callout A"), ShapeInfo("Callout B")]
veff = [Effect(0, MEDIA, trigger=AFTER, dur=10.0), Effect(1, FADE, trigger=AFTER, dur=0.5), Effect(2, FADE, trigger=CLICK, dur=0.5)]
av = T.analyze(veff, vshapes)
check("播放视频", av.media and av.clicks)
check("视频自己是一层（图片在页面上），不动", av.layers[0].media and not av.layers[0].animated and av.layers[0].shapes == [0])
check("两条提示都在视频之后", av.layers[1].after_media and av.layers[2].after_media)
check("提示 A 是第 1 步，提示 B 是第 2 步（视频的 10 秒不算进步骤里）",
      av.layers[1].enter[:2] == (0, 0.0) and av.layers[2].enter[:2] == (1, 0.0), (av.layers[1].enter, av.layers[2].enter))
check("每步说的内容", av.beat_texts(only_after_media=True) == ["Callout A", "Callout B"])
veff2 = [Effect(1, FADE, trigger=CLICK), Effect(0, MEDIA, trigger=AFTER, dur=10.0), Effect(2, FADE, trigger=WITH, delay=3.0, dur=0.5),
         Effect(2, FADE, trigger=AFTER, dur=0.5)]
av2 = T.analyze(veff2, vshapes)
check("视频之前出现的不算「视频之后」", not av2.layers[1].after_media)
veff3 = [Effect(0, MEDIA, trigger=CLICK, dur=10.0), Effect(1, FADE, trigger=WITH, delay=3.0, dur=0.5), Effect(2, FADE, trigger=WITH, delay=10.0, dur=0.5)]
av3 = T.analyze(veff3, vshapes)
check("视频播放当中出现的（第 3 秒）不是「视频之后」，播完时才出现的（第 10 秒）是",
      (not av3.layers[1].after_media) and av3.layers[2].after_media, (av3.layers[1].after_media, av3.layers[2].after_media))

print("\n== 8. 项目里存的层：分成「说的单位」 ==")
items = [RevealItem(text="bg"), RevealItem(text="card", beat=0), RevealItem(text="txt", beat=0, offset=0.0),
         RevealItem(text="p1", beat=1), RevealItem(text="late", beat=2, after_media=True), RevealItem(text="", beat=2, after_media=True)]
check("每一步是一个单位，静态层不算", T.units(items) == [[1, 2], [3], [4, 5]], T.units(items))
check("文字按步骤合起来", T.unit_texts(items) == ["card\ntxt", "p1", "late"], T.unit_texts(items))
check("只要视频之后的", T.units(items, only_after_media=True) == [[4, 5]])
legacy = [RevealItem(text="a"), RevealItem(text="b")]
check("按位置分组的老数据：每条一个单位", T.units(legacy) == [[0], [1]] and T.unit_texts(legacy) == ["a", "b"])


print("\n== 9. 擦除、飞入、瞬间出现、动作路径 ==")
WIPE_, CIRCLE = 22, 6
a = T.analyze([Effect(0, WIPE_, trigger=CLICK, direction="top"), Effect(1, 2, trigger=CLICK, direction="left"),
               Effect(2, CIRCLE, trigger=CLICK, dur=0.01), Effect(3, CIRCLE, trigger=CLICK, dur=0.25)],
              [ShapeInfo("w"), ShapeInfo("f"), ShapeInfo("c"), ShapeInfo("d")])
check("擦除、飞入记下是从哪一边来的", [(ly.enter[2], ly.side) for ly in a.layers[:2]] == [("wipe", "top"), ("fly", "left")],
      [(ly.enter[2], ly.side) for ly in a.layers[:2]])
check("时长很短（10 毫秒的「圆形扩展」光点）的当瞬间出现，不短的淡入并记下时长",
      a.layers[2].enter[2] == "appear" and a.layers[3].enter[2] == "fade" and abs(a.layers[3].dur - 0.25) < 1e-9,
      [(ly.enter[2], ly.dur) for ly in a.layers[2:]])
# the light spot of a real deck: it appears on a click, then "after previous" moves along a line for 2 seconds (a start and a slow end)
spot = [Effect(1, APPEAR, trigger=CLICK, dur=0.001),
        Effect(1, 120, trigger=AFTER, dur=2.0, cls="path", path="M -8.33333E-7 2.77556E-17 L 0.45 0.6 ", accel=0.24, decel=0.28)]
a = T.analyze(spot, [ShapeInfo("steps"), ShapeInfo("")])
spot_ly = a.layers[-1]
check("光点：先出现，同一步里接着沿路径移动 2 秒（不是新的一步）", a.beats == 1 and spot_ly.enter[2] == "appear" and spot_ly.motion is not None
      and spot_ly.motion[0] == 0 and abs(spot_ly.motion[2] - 2.0) < 1e-9 and abs(spot_ly.motion[4] - 0.24) < 1e-9, spot_ly)
check("路径交给渲染（相对幻灯片的比例）", "L 0.45 0.6" in spot_ly.motion[3])
meta = T.to_meta(a)[-1]
check("存进项目的数据里有路径、起止、缓入缓出", meta["motion_path"].startswith("M ") and meta["motion_dur"] == 2.0 and meta["motion_accel"] == 0.24
      and meta["motion_beat"] == 0 and meta["anim"] == "appear", meta)
a = T.analyze([Effect(0, FADE, trigger=CLICK), Effect(1, 120, trigger=CLICK, dur=1.0, cls="path", path="M 0 0 L 0.1 0 E")],
              [ShapeInfo("a"), ShapeInfo("moving picture")])
check("本身没有入场的东西只做动作路径：一直在，单独成层；这次点击不单独算一步，并进相邻的那一步（和只有退出的点击一样）",
      a.layers[1].motion is not None and not a.layers[1].animated and a.layers[1].motion[0] == 0 and a.beats == 1, a.layers)
a = T.analyze([Effect(0, FADE, trigger=CLICK)], [ShapeInfo("only one click")])
check("只有一次点击的页也按动画来（以前要两次以上）", a is not None and a.beats == 1)

print("\n== 10. 从 .pptx 文件里读擦除的方向和动作路径（PowerPoint 的接口不给） ==")
import tempfile  # noqa: E402

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

PNS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def timing_xml(effects):
    """effects = [(spid, presetID, presetClass, presetSubtype, nodeType, inner, accel, decel)] in the shape PowerPoint writes them."""
    out = []
    for n, (spid, pid, cls, sub, node, inner, accel, decel) in enumerate(effects):
        extra = (f' accel="{accel}"' if accel else "") + (f' decel="{decel}"' if decel else "")
        out.append(f'<p:par><p:cTn id="{10 + n}" presetID="{pid}" presetClass="{cls}" presetSubtype="{sub}"{extra} fill="hold" nodeType="{node}">'
                   f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{50 + n}" dur="1" fill="hold"/>'
                   f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:set>{inner}</p:childTnLst></p:cTn></p:par>')
    return (f'<p:timing xmlns:p="{PNS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold">'
            f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/>'
            f'</p:stCondLst><p:childTnLst>{"".join(out)}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
            f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


prs = Presentation()
sld = prs.slides.add_slide(prs.slide_layouts[6])
shp = [sld.shapes.add_shape(1, Inches(i), Inches(1), Inches(1), Inches(1)) for i in range(3)]
ids = [x.shape_id for x in shp]
wipe_in = '<p:animEffect transition="in" filter="wipe(up)"><p:cBhvr><p:cTn id="90" dur="500"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:animEffect>' % ids[0]
motion = '<p:animMotion origin="layout" path="M 0 0 L 0.45 0.6 " pathEditMode="fixed" rAng="0" ptsTypes="AA"><p:cBhvr><p:cTn id="91" dur="2000" fill="hold"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:animMotion>' % ids[1]
xml = timing_xml([(ids[0], 22, "entr", 1, "clickEffect", wipe_in, 0, 0), (ids[1], 1, "entr", 0, "clickEffect", "", 0, 0),
                  (ids[1], 42, "path", 0, "afterEffect", motion, 24000, 28000), (ids[2], 22, "entr", 8, "clickEffect", "", 0, 0)])
sld._element.append(etree.fromstring(xml))
deckf = Path(tempfile.gettempdir()) / "vt_timing_test.pptx"
prs.save(deckf)
xe = T.xml_effects(deckf, 1)
check("按顺序读出主序列里的每个动画：类别、目标、方向、路径、缓入缓出", [x["cls"] for x in xe] == ["entr", "entr", "path", "entr"]
      and [x["spid"] for x in xe] == [ids[0], ids[1], ids[1], ids[2]] and xe[0]["sub"] == 1 and xe[3]["sub"] == 8
      and xe[2]["path"].startswith("M 0 0 L 0.45") and abs(xe[2]["accel"] - 0.24) < 1e-9 and abs(xe[2]["decel"] - 0.28) < 1e-9, xe)
check("不存在的页：空", T.xml_effects(deckf, 9) == [] and T.xml_effects(Path("nope.pptx"), 1) == [])
effs = [Effect(0, WIPE_, shape_id=ids[0]), Effect(1, APPEAR, shape_id=ids[1]), Effect(1, 120, trigger=AFTER, shape_id=ids[1]), Effect(2, WIPE_, shape_id=ids[2])]
check("两份名单一一对上才补充", T.enrich(effs, xe))
check("补上了：擦除从上边（presetSubtype 1）、从左边（8）；路径和缓入缓出", effs[0].direction == "top" and effs[3].direction == "left"
      and effs[2].cls == "path" and effs[2].path.startswith("M 0 0") and abs(effs[2].accel - 0.24) < 1e-9, effs[2])
check("方向是按位记的：从角上飞入是两边相加（左上 9、右上 3、右下 6、左下 12，PowerPoint 实测）",
      [T.side_of(x) for x in (1, 2, 4, 8, 9, 3, 6, 12)] == ["top", "right", "bottom", "left", "top-left", "top-right", "bottom-right", "bottom-left"])
check("不是方向的数（其他效果的子类型、上下同时）不当方向", [T.side_of(x) for x in (0, 5, 10, 16, 21, 32)] == [""] * 6)
bad = [Effect(0, WIPE_, shape_id=ids[2])] + effs[1:]
check("目标对不上（名单不是同一个）：不乱补", not T.enrich(bad, xe) and bad[0].direction == "")
check("个数不同：不补", not T.enrich(effs[:3], xe))
a = T.analyze(effs, [ShapeInfo("a"), ShapeInfo("b"), ShapeInfo("c")])
check("补充后的分析：擦除的方向、光点的路径都带上了", a.layers[0].enter[2] == "wipe" and a.layers[0].side == "top"
      and a.layers[1].motion is not None and a.layers[2].side == "left", [(ly.enter, ly.side, ly.motion) for ly in a.layers])
deckf.unlink(missing_ok=True)

print("\n== 11. 从黑底、白底两次导出算出透明图层：只在有内容的范围里算，结果和整幅计算逐像素一样 ==")
import numpy as np  # noqa: E402
from PIL import Image, ImageDraw, ImageFilter  # noqa: E402


def matte_full(on_black, on_white):
    """The straightforward version over the whole picture (the reference)."""
    b = np.asarray(on_black.convert("RGB"), dtype=np.float32)
    w = np.asarray(on_white.convert("RGB"), dtype=np.float32)
    a = np.clip(1.0 - (w - b).mean(axis=2) / 255.0, 0.0, 1.0)
    a[a < 0.012] = 0.0
    a[a > 0.988] = 1.0
    ys, xs = np.nonzero(a)
    if not len(xs):
        return None
    x0, x1, y0, y1 = int(xs.min()), int(xs.max()) + 1, int(ys.min()), int(ys.max()) + 1
    a, b = a[y0:y1, x0:x1], b[y0:y1, x0:x1]
    rgb = np.clip(b / np.maximum(a[..., None], 1e-3), 0, 255)
    return Image.fromarray(np.dstack([rgb, a * 255.0 + 0.5]).astype(np.uint8), "RGBA"), x0, y0


def exports(draw_fn, size=(320, 180)):
    """What PowerPoint would export: the same element over black and over white."""
    layer = Image.new("RGBA", size, (0, 0, 0, 0))
    draw_fn(ImageDraw.Draw(layer), layer)
    out = []
    for bg in ((0, 0, 0), (255, 255, 255)):
        im = Image.new("RGB", size, bg)
        im.paste(layer, (0, 0), layer)
        out.append(im)
    return out


cases = {
    "实心卡片": lambda d, im: d.rectangle((40, 30, 140, 90), fill=(40, 120, 200, 255)),
    "半透明卡片": lambda d, im: d.rectangle((60, 40, 200, 120), fill=(200, 40, 40, 128)),
    "黑色的字在黑底上看不见（白底上看得见）": lambda d, im: d.text((20, 20), "Black text", fill=(0, 0, 0, 255)),
    "白色的字在白底上看不见": lambda d, im: d.text((150, 100), "White text", fill=(255, 255, 255, 255)),
    "贴着边的元素": lambda d, im: d.rectangle((0, 0, 30, 179), fill=(10, 200, 10, 255)),
    "柔和的阴影（大片很淡的像素）": lambda d, im: (d.ellipse((80, 40, 240, 140), fill=(0, 0, 0, 90)),
                                        im.paste(im.filter(ImageFilter.GaussianBlur(12)))),
    "两块分开的元素（一个图层里两个形状）": lambda d, im: (d.rectangle((10, 10, 40, 40), fill=(255, 128, 0, 255)),
                                         d.rectangle((260, 130, 300, 170), fill=(0, 128, 255, 200))),
}
for name, fn in cases.items():
    ib, iw = exports(fn)
    got, ref = T.matte(ib, iw), matte_full(ib, iw)
    same = (got is None and ref is None) or (got is not None and ref is not None and got[1:] == ref[1:]
                                             and got[0].size == ref[0].size and got[0].tobytes() == ref[0].tobytes())
    check(f"{name}：和整幅计算完全一样", same, (got and (got[1:], got[0].size), ref and (ref[1:], ref[0].size)))
ib, iw = exports(lambda d, im: None)
check("什么都没有：None", T.matte(ib, iw) is None and matte_full(ib, iw) is None)

print("\n" + ("全部通过" if not fails else f"失败 {len(fails)} 项：{fails}"))
